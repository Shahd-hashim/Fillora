"""Read an uploaded CV (PDF, DOCX or TXT) and clean up the profile the LLM extracted from it."""
import io

from docx import Document
from pypdf import PdfReader


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(lines).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")


def _s(v) -> str:
    return v.strip() if isinstance(v, str) else ("" if v is None else str(v))


def _joined(v) -> str:
    return ", ".join(x for x in (_s(i) for i in v) if x) if isinstance(v, list) else _s(v)


def normalize_profile(raw: dict) -> dict:
    """Same shape the extension stores: plain strings, comma separated skills, clean lists."""
    def entries(key: str, fields: list[str]) -> list[dict]:
        items = raw.get(key)
        return [
            {f: _s(e.get(f)) for f in fields}
            for e in (items if isinstance(items, list) else [])
            if isinstance(e, dict)
        ]

    return {
        "fullName": _s(raw.get("fullName")),
        "email": _s(raw.get("email")),
        "phone": _s(raw.get("phone")),
        "location": _s(raw.get("location")),
        "linkedin": _s(raw.get("linkedin")),
        "github": _s(raw.get("github")),
        "website": _s(raw.get("website")),
        "summary": _s(raw.get("summary")),
        "skills": _joined(raw.get("skills")),
        "languages": _joined(raw.get("languages")),
        "education": entries("education", ["school", "degree", "field", "start", "end"]),
        "experience": entries("experience", ["company", "title", "start", "end", "description"]),
    }
